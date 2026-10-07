"""Create a concise Word report from outputs/metrics.json."""
from pathlib import Path
import json, argparse
from docx import Document
from docx.shared import Inches

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--metrics', default='outputs/metrics.json'); ap.add_argument('--out', default='HW2aCNNautoencoder_report.docx'); args=ap.parse_args()
    m=json.loads(Path(args.metrics).read_text()); d=Document(); d.add_heading('Autoencoder Anomaly Detection for Document Triage', 0)
    d.add_paragraph('This report uses Fashion-MNIST as a reproducible proxy for invoice thumbnails. Classes 0–5 are treated as routine documents during training; classes 6–9 are held out as suspicious documents during evaluation.')
    d.add_heading('Method', 1); c=m['config']; d.add_paragraph(f"The experiment used seed {c['seed']}, {c['epochs']} epochs, batch size {c['batch_size']}, MSE reconstruction loss, Adam with learning rate 1e-3, and a {c['percentile']}th-percentile threshold estimated from normal validation errors. The convolutional autoencoder uses strided convolutional blocks and a {c['bottleneck']}-channel spatial bottleneck. The dense baseline uses fully connected layers and a {c['latent_dim']}-dimensional bottleneck.")
    d.add_heading('Results', 1); t=d.add_table(rows=1, cols=6); t.style='Light Shading Accent 1'
    for cell, txt in zip(t.rows[0].cells, ['Model','Threshold','Precision','Recall','ROC AUC','Confusion matrix']): cell.text=txt
    for name in ['conv','dense']:
        r=m[name]; cells=t.add_row().cells
        vals=[name, f"{r['threshold']:.6f}", f"{r['precision']:.3f}", f"{r['recall']:.3f}", f"{r['roc_auc']:.3f}", str(r['confusion_matrix'])]
        for cell, txt in zip(cells, vals): cell.text=txt
    d.add_heading('Interpretation', 1); d.add_paragraph('A reconstruction-error detector assumes that the training distribution is sufficiently representative of routine documents. A high error means the model did not reproduce the input well, which can indicate unfamiliar visual structure. The convolutional model is preferred because convolutions preserve local spatial relationships and share parameters across the thumbnail; the dense model discards this inductive bias when flattening the image.')
    d.add_heading('Limitations and next steps', 1); d.add_paragraph('Fashion-MNIST is not a real invoice dataset, so the anomaly results should be interpreted as a pipeline demonstration rather than evidence of production invoice-screening accuracy. A real deployment would require document-specific training data, a validation protocol that reflects changing vendors and templates, threshold calibration for alert volume, and monitoring for concept drift. Denoising training and latent-space visualization are useful extensions.')
    for p in ['conv_loss.png','conv_normal_examples.png','conv_anomaly_examples.png','conv_error_distribution.png']:
        q=Path(args.metrics).parent/p
        if q.exists():
            d.add_picture(str(q), width=Inches(6.2)); d.add_paragraph(p)
    d.save(args.out)

if __name__=='__main__': main()
